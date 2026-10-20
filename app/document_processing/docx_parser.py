"""Extracción de texto estructurado desde archivos .docx."""

import logging
from pathlib import Path

from docx import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.table import Table
from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)

# Mapeo de estilos de encabezado a niveles markdown
_HEADING_STYLES: dict[str, str] = {
    "Heading 1": "# ",
    "Heading 2": "## ",
    "Heading 3": "### ",
    "Heading 4": "#### ",
    "Heading 5": "##### ",
    "Heading 6": "###### ",
    "Title": "# ",
    "Subtitle": "## ",
    # Variantes en español
    "Título": "# ",
    "Título 1": "# ",
    "Título 2": "## ",
    "Título 3": "### ",
    "Título 4": "#### ",
    "Subtítulo": "## ",
}


def _paragraph_to_text(paragraph: Paragraph) -> str:
    """Convierte un párrafo a texto con formato markdown para encabezados.

    Args:
        paragraph: Párrafo de python-docx.

    Returns:
        Texto del párrafo con prefijo markdown si es encabezado.
    """
    text = paragraph.text.strip()
    if not text:
        return ""

    style_name = paragraph.style.name if paragraph.style else ""

    # Verificar si es un encabezado por nombre de estilo
    prefix = _HEADING_STYLES.get(style_name, "")

    # Verificar también por nombre base con número (e.g., "Heading 1")
    if not prefix and style_name:
        for heading_key, heading_prefix in _HEADING_STYLES.items():
            if style_name.lower().startswith(heading_key.lower()):
                prefix = heading_prefix
                break

    return f"{prefix}{text}"


def _table_to_text(table: Table) -> str:
    """Convierte una tabla a formato de texto con separadores de pipe.

    Args:
        table: Tabla de python-docx.

    Returns:
        Tabla formateada como texto con pipes (estilo markdown).
    """
    rows_text: list[str] = []

    for row_idx, row in enumerate(table.rows):
        cells = [cell.text.strip().replace("|", "\\|") for cell in row.cells]
        row_line = "| " + " | ".join(cells) + " |"
        rows_text.append(row_line)

        # Agregar separador después de la primera fila (encabezado)
        if row_idx == 0:
            separator = "| " + " | ".join("---" for _ in cells) + " |"
            rows_text.append(separator)

    return "\n".join(rows_text)


def _extract_header_footer_text(doc: DocxDocument) -> tuple[str, str]:
    """Extrae texto de encabezados y pies de página del documento.

    Args:
        doc: Documento python-docx.

    Returns:
        Tupla con (texto_encabezados, texto_pies).
    """
    headers_text: list[str] = []
    footers_text: list[str] = []

    try:
        for section in doc.sections:
            # Encabezados
            header = section.header
            if header and not header.is_linked_to_previous:
                for paragraph in header.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        headers_text.append(text)

            # Pies de página
            footer = section.footer
            if footer and not footer.is_linked_to_previous:
                for paragraph in footer.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        footers_text.append(text)
    except Exception as e:
        logger.warning("Error al extraer encabezados/pies de página: %s", e)

    return "\n".join(headers_text), "\n".join(footers_text)


def extract_from_docx(file_path: str | Path) -> str:
    """Extrae texto estructurado de un archivo .docx.

    Procesa párrafos, tablas, encabezados y pies de página preservando
    la estructura del documento.

    Args:
        file_path: Ruta al archivo .docx.

    Returns:
        Texto extraído con estructura preservada en formato markdown.

    Raises:
        FileNotFoundError: Si el archivo no existe.
        ValueError: Si el archivo no es un .docx válido.
    """
    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"Archivo no encontrado: {file_path}")

    if file_path.suffix.lower() != ".docx":
        raise ValueError(f"El archivo no es un .docx válido: {file_path.name}")

    try:
        doc = DocxDocument(str(file_path))
    except Exception as e:
        raise ValueError(f"No se pudo leer el archivo .docx: {e}") from e

    partes: list[str] = []

    # ── Extraer encabezados y pies de página ──
    header_text, footer_text = _extract_header_footer_text(doc)

    if header_text:
        partes.append(f"[ENCABEZADO]\n{header_text}\n[/ENCABEZADO]")

    # ── Extraer cuerpo del documento ──
    # Iterar sobre los elementos del cuerpo en orden para mantener
    # la posición relativa de párrafos y tablas
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

        if tag == "p":
            # Es un párrafo
            paragraph = Paragraph(child, doc)
            text = _paragraph_to_text(paragraph)
            if text:
                partes.append(text)

        elif tag == "tbl":
            # Es una tabla
            table = Table(child, doc)
            table_text = _table_to_text(table)
            if table_text:
                partes.append(f"\n{table_text}\n")

    # ── Agregar pies de página al final ──
    if footer_text:
        partes.append(f"\n[PIE]\n{footer_text}\n[/PIE]")

    resultado = "\n".join(partes)

    logger.info(
        "Texto extraído de '%s': %d caracteres, %d partes",
        file_path.name,
        len(resultado),
        len(partes),
    )

    return resultado
