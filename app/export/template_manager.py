"""Gestor de plantillas institucionales de Word (.docx) con docxtpl."""

import io
import logging
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docxtpl import DocxTemplate

from app.config import get_settings
from app.export.docx_builder import PlanDocxBuilder
from app.export.styles import (
    COLOR_PRIMARY,
    COLOR_SECONDARY,
    HEX_ALT_ROW_BG,
    HEX_HEADER_BG,
    MARGIN_BOTTOM,
    MARGIN_LEFT,
    MARGIN_RIGHT,
    MARGIN_TOP,
    set_cell_margins,
    set_cell_shading,
)
from app.schemas.plan import PlanDiario

logger = logging.getLogger(__name__)
settings = get_settings()

TEMPLATES_DIR = Path(settings.upload_dir) / "templates"
BASE_TEMPLATE_PATH = TEMPLATES_DIR / "plantilla_modelo_kantan.docx"


def ensure_templates_dir():
    """Asegura que el directorio de plantillas exista."""
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)


def create_base_model_template() -> Path:
    """Crea una plantilla modelo .docx con etiquetas Jinja si no existe aún.

    Esta plantilla sirve de base descargable para que el docente pueda abrirla
    en Word, pegar el logo de su escuela en la cabecera, guardarla y subirla.
    """
    ensure_templates_dir()
    if BASE_TEMPLATE_PATH.exists():
        return BASE_TEMPLATE_PATH

    doc = Document()
    for section in doc.sections:
        section.top_margin = MARGIN_TOP
        section.bottom_margin = MARGIN_BOTTOM
        section.left_margin = MARGIN_LEFT
        section.right_margin = MARGIN_RIGHT

    # Cabecera / Banner Institucional
    head_table = doc.add_table(rows=1, cols=2)
    head_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    head_table.autofit = False

    # Columna izquierda: Espacio para Logo
    cell_logo = head_table.rows[0].cells[0]
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_logo = p_logo.add_run("🖼️ [PEGUE AQUÍ EL LOGO O ESCUDO DE SU ESCUELA]")
    r_logo.font.size = Pt(9)
    r_logo.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
    r_logo.italic = True
    set_cell_margins(cell_logo, top=120, bottom=120, left=150, right=150)
    set_cell_shading(cell_logo, "F8F9FA")

    # Columna derecha: Título y Escuela
    cell_title = head_table.rows[0].cells[1]
    p_t = cell_title.paragraphs[0]
    p_t.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_t1 = p_t.add_run("PLAN DIARIO DE CLASE\n")
    r_t1.bold = True
    r_t1.font.size = Pt(15)
    r_t1.font.color.rgb = COLOR_PRIMARY

    r_t2 = p_t.add_run("{{ INSTITUCION }}\n")
    r_t2.font.size = Pt(11)
    r_t2.bold = True
    r_t2.font.color.rgb = COLOR_SECONDARY

    r_t3 = p_t.add_run("Ciclo Lectivo {{ FECHA }}")
    r_t3.font.size = Pt(9.5)
    r_t3.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
    set_cell_margins(cell_title, top=120, bottom=120, left=150, right=150)

    # Espacio
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Tabla de Datos Generales
    datos_table = doc.add_table(rows=3, cols=2)
    datos_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    info_fields = [
        ("Materia:", "{{ MATERIA }}", "Grado / Nivel:", "{{ GRADO }}"),
        ("Docente:", "{{ DOCENTE }}", "Fecha:", "{{ FECHA }}"),
        ("Duración estimada:", "{{ DURACION }} minutos", "Modalidad:", "Presencial"),
    ]

    for i, (l_lbl, l_val, r_lbl, r_val) in enumerate(info_fields):
        row = datos_table.rows[i]
        c0 = row.cells[0]
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(f"{l_lbl} ")
        r0.bold = True
        p0.add_run(l_val)

        c1 = row.cells[1]
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(f"{r_lbl} ")
        r1.bold = True
        p1.add_run(r_val)

        for cell in (c0, c1):
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            if i % 2 == 0:
                set_cell_shading(cell, HEX_ALT_ROW_BG)

    # Secciones
    def add_sec_title(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"■  {title_text.upper()}")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_PRIMARY

    # Objetivo
    add_sec_title("1. Objetivo de Aprendizaje")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(6)
    p_obj.add_run("{{ OBJETIVO }}")

    # Contenidos
    add_sec_title("2. Contenidos y Saberes")
    p_cont = doc.add_paragraph()
    p_cont.paragraph_format.space_after = Pt(6)
    p_cont.add_run("{{ CONTENIDOS }}")

    # Secuencia Didáctica
    add_sec_title("3. Secuencia Didáctica")
    sec_table = doc.add_table(rows=4, cols=3)
    sec_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Momento de la Clase", "Duración", "Actividades y Consignas"]
    for j, h_text in enumerate(headers):
        c = sec_table.rows[0].cells[j]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        set_cell_shading(c, HEX_HEADER_BG)
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)

    fases = [
        ("Inicio (Apertura)", "{{ INICIO_DURACION }} min", "{{ INICIO_DESC }}\n{{ INICIO_ACTIVIDADES }}"),
        ("Desarrollo (Construcción)", "{{ DESARROLLO_DURACION }} min", "{{ DESARROLLO_DESC }}\n{{ DESARROLLO_ACTIVIDADES }}"),
        ("Cierre (Evaluación/Síntesis)", "{{ CIERRE_DURACION }} min", "{{ CIERRE_DESC }}\n{{ CIERRE_ACTIVIDADES }}"),
    ]

    for idx, (fase_nom, fase_dur, fase_desc) in enumerate(fases, start=1):
        row = sec_table.rows[idx]
        row.cells[0].paragraphs[0].add_run(fase_nom).bold = True
        row.cells[1].paragraphs[0].add_run(fase_dur)
        row.cells[2].paragraphs[0].add_run(fase_desc)
        for cell in row.cells:
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            if idx % 2 == 1:
                set_cell_shading(cell, HEX_ALT_ROW_BG)

    # Recursos
    add_sec_title("4. Recursos y Materiales")
    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.space_after = Pt(6)
    p_rec.add_run("{{ RECURSOS }}")

    # Evaluación
    add_sec_title("5. Evaluación")
    p_eval = doc.add_paragraph()
    p_eval.paragraph_format.space_after = Pt(4)
    r_et = p_eval.add_run("Tipo: ")
    r_et.bold = True
    p_eval.add_run("{{ EVALUACION_TIPO }}\n")
    p_eval.add_run("{{ EVALUACION_DESC }}\n\nCriterios de logro:\n{{ EVALUACION_CRITERIOS }}")

    doc.save(str(BASE_TEMPLATE_PATH))
    logger.info("Plantilla modelo base creada en %s", BASE_TEMPLATE_PATH)
    return BASE_TEMPLATE_PATH


def get_user_template_path(user_id: int) -> Path | None:
    """Retorna la ruta de la plantilla personalizada del docente si existe."""
    ensure_templates_dir()
    custom_path = TEMPLATES_DIR / f"user_{user_id}_template.docx"
    if custom_path.exists():
        return custom_path
    return None


def save_user_template(user_id: int, file_bytes: bytes, original_filename: str) -> Path:
    """Guarda la plantilla personalizada subida por el docente."""
    ensure_templates_dir()
    target_path = TEMPLATES_DIR / f"user_{user_id}_template.docx"
    with open(target_path, "wb") as f:
        f.write(file_bytes)

    # Guardar metadata del nombre original
    meta_path = TEMPLATES_DIR / f"user_{user_id}_template.meta"
    meta_path.write_text(original_filename, encoding="utf-8")

    return target_path


def remove_user_template(user_id: int) -> bool:
    """Elimina la plantilla personalizada y vuelve a la predeterminada."""
    custom_path = get_user_template_path(user_id)
    if custom_path and custom_path.exists():
        custom_path.unlink()
        meta_path = TEMPLATES_DIR / f"user_{user_id}_template.meta"
        if meta_path.exists():
            meta_path.unlink()
        return True
    return False


def get_user_template_name(user_id: int) -> str | None:
    """Retorna el nombre del archivo de la plantilla personalizada activa."""
    meta_path = TEMPLATES_DIR / f"user_{user_id}_template.meta"
    if meta_path.exists():
        return meta_path.read_text(encoding="utf-8").strip()
    custom_path = get_user_template_path(user_id)
    if custom_path:
        return "Plantilla_Institucional.docx"
    return None


def build_template_context(plan: PlanDiario) -> dict:
    """Construye el diccionario de variables Jinja2 para rellenar la plantilla."""
    dg = plan.datos_generales
    sd = plan.secuencia_didactica
    ev = plan.evaluacion

    cont_str = "\n".join(f"• {c}" for c in plan.contenidos)
    rec_str = "\n".join(f"• {r}" for r in plan.recursos)

    crit_str = ""
    if ev.criterios:
        crit_str = "\n".join(f"• {c}" for c in ev.criterios)

    ini_act = "\n".join(f"• {a}" for a in sd.inicio.actividades) if sd.inicio.actividades else ""
    des_act = "\n".join(f"• {a}" for a in sd.desarrollo.actividades) if sd.desarrollo.actividades else ""
    cie_act = "\n".join(f"• {a}" for a in sd.cierre.actividades) if sd.cierre.actividades else ""

    ctx = {
        # Generales
        "MATERIA": dg.materia,
        "materia": dg.materia,
        "GRADO": dg.grado,
        "grado": dg.grado,
        "FECHA": str(dg.fecha),
        "fecha": str(dg.fecha),
        "DOCENTE": dg.docente or "No especificado",
        "docente": dg.docente or "No especificado",
        "DURACION": dg.duracion_minutos,
        "duracion": dg.duracion_minutos,
        "INSTITUCION": dg.institucion or "Institución Educativa",
        "institucion": dg.institucion or "Institución Educativa",

        # Objetivo y Contenidos
        "OBJETIVO": plan.objetivo,
        "objetivo": plan.objetivo,
        "CONTENIDOS": cont_str,
        "contenidos": cont_str,

        # Secuencia
        "INICIO_DURACION": sd.inicio.duracion_minutos,
        "inicio_duracion": sd.inicio.duracion_minutos,
        "INICIO_DESC": sd.inicio.descripcion,
        "inicio_desc": sd.inicio.descripcion,
        "INICIO_ACTIVIDADES": ini_act,
        "inicio_actividades": ini_act,

        "DESARROLLO_DURACION": sd.desarrollo.duracion_minutos,
        "desarrollo_duracion": sd.desarrollo.duracion_minutos,
        "DESARROLLO_DESC": sd.desarrollo.descripcion,
        "desarrollo_desc": sd.desarrollo.descripcion,
        "DESARROLLO_ACTIVIDADES": des_act,
        "desarrollo_actividades": des_act,

        "CIERRE_DURACION": sd.cierre.duracion_minutos,
        "cierre_duracion": sd.cierre.duracion_minutos,
        "CIERRE_DESC": sd.cierre.descripcion,
        "cierre_desc": sd.cierre.descripcion,
        "CIERRE_ACTIVIDADES": cie_act,
        "cierre_actividades": cie_act,

        # Recursos y Evaluación
        "RECURSOS": rec_str,
        "recursos": rec_str,
        "EVALUACION_TIPO": ev.tipo,
        "evaluacion_tipo": ev.tipo,
        "EVALUACION_DESC": ev.descripcion,
        "evaluacion_desc": ev.descripcion,
        "EVALUACION_CRITERIOS": crit_str,
        "evaluacion_criterios": crit_str,
    }
    return ctx


def render_plan_to_docx(plan: PlanDiario, user_id: int = 1) -> io.BytesIO:
    """Renderiza el plan utilizando la plantilla personalizada o el builder predeterminado.

    Si el usuario subió una plantilla personalizada, docxtpl la rellena manteniendo
    su logo, colores, membrete e imágenes intactos.
    Si no tiene plantilla o si la plantilla personalizada falla, usa PlanDocxBuilder como respaldo seguro.
    """
    user_tpl_path = get_user_template_path(user_id)

    if user_tpl_path and user_tpl_path.exists():
        try:
            logger.info("Renderizando plan con plantilla personalizada de usuario %d: %s", user_id, user_tpl_path)
            tpl = DocxTemplate(str(user_tpl_path))
            context = build_template_context(plan)
            tpl.render(context)

            buffer = io.BytesIO()
            tpl.save(buffer)
            buffer.seek(0)
            return buffer
        except Exception as e:
            logger.warning("Error al renderizar plantilla personalizada (%s). Usando generador estándar: %s", user_tpl_path, e)

    # Si no hay plantilla personalizada o hubo error, usar el generador estándar con estilo profesional
    builder = PlanDocxBuilder()
    return builder.build(plan)
