"""Pruebas para el gestor de plantillas institucionales de Word."""

import pytest
from app.export.template_manager import (
    create_base_model_template,
    render_plan_to_docx,
    save_user_template,
    remove_user_template,
)
from app.schemas.plan import (
    ActividadSecuencia,
    DatosGenerales,
    Evaluacion,
    PlanDiario,
    ProblemaPractico,
    SecuenciaDidactica,
)


@pytest.fixture
def sample_plan():
    return PlanDiario(
        datos_generales=DatosGenerales(
            materia="Matemáticas",
            grado="4° Grado",
            fecha="20/10/2026",
            duracion_minutos=50,
            docente="Profesor Test",
            institucion="Colegio San Martín",
        ),
        objetivo="Comprender el concepto de fracciones equivalentes.",
        contenidos=["Fracciones propias", "Fracciones impropias"],
        secuencia_didactica=SecuenciaDidactica(
            inicio=ActividadSecuencia(
                duracion_minutos=10,
                descripcion="Activación de saberes previos.",
                actividades=["Pregunta disparadora"],
            ),
            desarrollo=ActividadSecuencia(
                duracion_minutos=30,
                descripcion="Resolución de problemas con material concreto.",
                actividades=["Trabajo en parejas"],
            ),
            cierre=ActividadSecuencia(
                duracion_minutos=10,
                descripcion="Puesta en común y cierre.",
                actividades=["Ticket de salida"],
            ),
        ),
        recursos=["Pizarrón", "Fichas"],
        evaluacion=Evaluacion(
            tipo="Formativa",
            descripcion="Observación del trabajo en clase.",
            criterios=["Participación", "Resolución correcta"],
        ),
    )


def test_create_base_model_template():
    """Verifica que la plantilla base se cree físicamente en disco."""
    path = create_base_model_template()
    assert path.exists()
    assert path.suffix == ".docx"


def test_render_plan_default(sample_plan):
    """Verifica que el generador produzca un buffer BytesIO válido sin plantilla custom."""
    buffer = render_plan_to_docx(sample_plan, user_id=999)
    assert buffer is not None
    assert buffer.getbuffer().nbytes > 1000


def test_render_plan_with_custom_template(sample_plan):
    """Verifica que renderice exitosamente sobre una plantilla con etiquetas docxtpl."""
    base_path = create_base_model_template()
    with open(base_path, "rb") as f:
        file_bytes = f.read()

    save_user_template(user_id=888, file_bytes=file_bytes, original_filename="Mi_Escuela.docx")
    
    buffer = render_plan_to_docx(sample_plan, user_id=888)
    assert buffer is not None
    assert buffer.getbuffer().nbytes > 1000

    # Limpieza
    remove_user_template(user_id=888)


def test_render_plan_mec_structure():
    """Verifica la generación del formato oficial MEC Paraguay y ejercitario."""
    mec_plan = PlanDiario(
        datos_generales=DatosGenerales(
            materia="Matemática",
            grado="5° Grado",
            fecha="22/10/2026",
            mes="Octubre",
            ano=2026,
            duracion_minutos=80,
            docente="Prof. Marcelo",
            institucion="Escuela Básica N° 123",
        ),
        tema="Multiplicación y División de Fracciones",
        unidad_tematica="El número y las operaciones",
        capacidades=[
            "Comprende el enunciado del problema.",
            "Aplica algoritmos de multiplicación de fracciones.",
        ],
        indicadores=[
            "Identifica los datos y la incógnita.",
            "Simplifica antes de operar.",
            "Formula la respuesta con unidad correcta.",
        ],
        momentos_didacticos_resumen="Inicio con situación disparadora. Desarrollo de ejercicios en cuaderno. Cierre con R.S.A.",
        instrumento_evaluativo="R.S.A.",
        situacion_disparadora="Doña Juana tiene 3/4 kg de harina y necesita preparar 2 recetas...",
        conceptualizacion="Para multiplicar fracciones, se multiplican numeradores entre sí y denominadores entre sí.",
        problemas=[
            ProblemaPractico(
                numero=1,
                enunciado="En una granja hay 120 animales. 2/3 son aves. ¿Cuántas aves hay?",
                datos="Total: 120 animales. 2/3 son aves.",
                formula_o_dibujo="120 * (2/3)",
                solucion="120 * 2 / 3 = 80",
                respuesta="Hay 80 aves en la granja.",
            )
        ],
        ejercicios_aplicacion=["Calcula 3/4 * 2/5", "Resuelve 5/6 : 1/2"],
        objetivo="Resolver problemas de la vida cotidiana que involucren operaciones con fracciones.",
        contenidos=["Fracciones", "Multiplicación", "División"],
        secuencia_didactica=SecuenciaDidactica(
            inicio=ActividadSecuencia(duracion_minutos=15, descripcion="Activación", actividades=["Pregunta motivadora"]),
            desarrollo=ActividadSecuencia(duracion_minutos=50, descripcion="Práctica", actividades=["Ejercicios"]),
            cierre=ActividadSecuencia(duracion_minutos=15, descripcion="Evaluación", actividades=["R.S.A."]),
        ),
        recursos=["Cuaderno", "Pizarra"],
        evaluacion=Evaluacion(tipo="Formativa", descripcion="R.S.A."),
    )

    buffer = render_plan_to_docx(mec_plan, user_id=777)
    assert buffer is not None
    assert buffer.getbuffer().nbytes > 2000


def test_render_raw_template_without_jinja(sample_plan):
    """Verifica que un archivo .docx escolar real (sin tags Jinja2) conserve membrete e inserte el plan."""
    from docx import Document
    import io

    # Crear un docx institucional simulado sin etiquetas Jinja
    doc = Document()
    doc.add_heading("ESCUELA BÁSICA NACIONAL N° 45", level=1)
    doc.add_paragraph("Docente: Anterior")
    doc.add_paragraph("Materia: Matemática")
    doc.add_paragraph("Grado: 4°")
    doc.add_paragraph("Texto antiguo que debe ser reemplazado...")

    raw_io = io.BytesIO()
    doc.save(raw_io)
    raw_bytes = raw_io.getvalue()

    save_user_template(user_id=666, file_bytes=raw_bytes, original_filename="Colegio_Membrete.docx")

    buffer = render_plan_to_docx(sample_plan, user_id=666)
    assert buffer is not None
    assert buffer.getbuffer().nbytes > 1500

    # Limpieza
    remove_user_template(user_id=666)

